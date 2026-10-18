import copy
from pptx import Presentation
from pptx.util import Inches as I, Pt, Emu
from pptx.dml.color import RGBColor as C
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY=C(0x0F,0x2A,0x3D); TEAL=C(0x0B,0x7F,0x88); CYAN=C(0x5D,0xE0,0xE6)
RED=C(0xC0,0x39,0x2B); WHITE=C(0xFF,0xFF,0xFF); PALE=C(0xF7,0xFB,0xFC); LINE=C(0xA9,0xD3,0xD8)
MUTE=C(0x4A,0x63,0x72); AMBER=C(0xE6,0x9B,0x1E); GREEN=C(0x1E,0x8E,0x5A)
HEAD="Times"; BODY="Arial"

p=Presentation('template.pptx'); S=list(p.slides)

def box(s,x,y,w,h,fill=PALE,line=LINE,shape=MSO_SHAPE.ROUNDED_RECTANGLE,r=0.08,lw=1.5):
    sh=s.shapes.add_shape(shape,I(x),I(y),I(w),I(h))
    if shape==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=r
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False
    return sh

def txt(s,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP,align=PP_ALIGN.LEFT,margin=0.0):
    """paras: list of (text,size,bold,color,font[,space_after]) ; text may be list of runs"""
    tb=s.shapes.add_textbox(I(x),I(y),I(w),I(h)); tf=tb.text_frame; tf.word_wrap=True
    tf.vertical_anchor=anchor
    for a in ('margin_left','margin_right','margin_top','margin_bottom'): setattr(tf,a,I(margin))
    for i,pa in enumerate(paras):
        t,sz,b,col,font=pa[:5]; sp=pa[5] if len(pa)>5 else 6
        para=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        para.alignment=align; para.space_after=Pt(sp)
        for rt in (t if isinstance(t,list) else [(t,b,col)]):
            r=para.add_run(); r.text=rt[0]; f=r.font; f.size=Pt(sz); f.bold=rt[1]; f.color.rgb=rt[2]; f.name=font
    return tb

def card(s,x,y,w,h,title,body,accent=TEAL,tsz=30,bsz=21,fill=PALE):
    box(s,x,y,w,h,fill=fill)
    bar=box(s,x,y+0.25,0.1,0.5,fill=accent,line=None,shape=MSO_SHAPE.RECTANGLE)
    paras=[(title,tsz,True,NAVY,HEAD,8)]+[(b,bsz,False,MUTE if i else NAVY,BODY,7) for i,b in enumerate(body)]
    txt(s,x+0.35,y+0.2,w-0.65,h-0.3,paras)

def arrow(s,x1,y1,x2,y2,col=TEAL,w=3,head=True):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,I(x1),I(y1),I(x2),I(y2))
    c.line.color.rgb=col; c.line.width=Pt(w)
    if head:
        ln=c.line._get_or_add_ln(); te=etree.SubElement(ln,qn('a:tailEnd')); te.set('type','triangle'); te.set('w','lg'); te.set('len','lg')
    return c

def set_title(s,text,size=None):
    for sh in s.shapes:
        if sh.has_text_frame and sh.name.startswith('Google Shape') and sh.top<I(1.2) and sh.width>I(4) and sh.text_frame.text.strip():
            r=sh.text_frame.paragraphs[0].runs[0]; r.text=text
            for extra in sh.text_frame.paragraphs[0].runs[1:]: extra.text=''
            sh.text_frame.word_wrap=True; sh.width=Emu(16200000)
            if size: r.font.size=Pt(size)
            return sh

# ---------- 1 title ----------
s=S[0]
for sh in s.shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith('YOUR NAME'):
        sh.width=Emu(9500000)
        vals={0:"Nikhil, Shreyas, Koustav, Deekshitha",2:"Chaukanna",4:" PS22 · AI Automation with n8n"}
        for i,v in vals.items():
            para=sh.text_frame.paragraphs[i]
            r=copy.deepcopy(para.runs[0]._r); para._p.insert(list(para._p).index(para.runs[0]._r)+1,r)
            rr=para.runs[1]; rr.text=v; rr.font.bold=False; rr.font.color.rgb=TEAL; rr.font.name=BODY; rr.font.size=Pt(22)
        for pa in sh.text_frame.paragraphs: pa.alignment=PP_ALIGN.LEFT

# ---------- 2 problem + solution ----------
s=S[1]
txt(s,8.2,1.75,10.6,1.0,[("Chaukanna (चौकन्ना): alert, watchful. The call guardian that acts for you.",26,False,NAVY,HEAD)],anchor=MSO_ANCHOR.MIDDLE)
box(s,1.125,3.25,6.5,7.2,fill=WHITE,line=RED,lw=2.5)
txt(s,1.5,3.5,5.8,6.8,[
 ("THE PROBLEM",20,True,RED,BODY,4),
 ("Digital-arrest & KYC scam calls run on panic",32,True,NAVY,HEAD,14),
 ("A fake police, CBI or bank caller keeps the victim on the line and demands secrecy, OTPs and transfers.",23,False,MUTE,BODY,12),
 ("The response is repetitive and fragmented: call the bank, call 1930, file on cybercrime.gov.in, warn family, keep evidence.",23,False,MUTE,BODY,12),
 ("In a panic, almost nobody does any of it in time, and scripts change language by language, so static block-lists miss them.",23,False,MUTE,BODY,6)])
txt(s,8.2,3.15,10.6,0.6,[("OUR SOLUTION: LISTEN → UNDERSTAND → INTERVENE → RESPOND",21,True,TEAL,BODY)])
steps=[("1","Listen","Calls run through the app, so both legs are heard separately and live."),
       ("2","Understand","Indic speech-to-text, then a typed verdict: scam type, pressure 0–3, OTP/PIN ask, secrecy demand."),
       ("3","Intervene","Only the victim hears a whispered warning in their own language, plus an alarm."),
       ("4","Respond (n8n)","An AI agent alerts family, drafts the 1930 complaint, logs evidence and updates a shared scam-number list.")]
for i,(n,t,b) in enumerate(steps):
    x=8.2+(i%2)*5.4; y=3.85+(i//2)*3.3
    box(s,x,y,5.2,3.1,fill=WHITE)
    c=box(s,x+0.3,y+0.3,0.75,0.75,fill=TEAL if i<3 else NAVY,line=None,shape=MSO_SHAPE.OVAL)
    txt(s,x+0.3,y+0.3,0.75,0.75,[(n,26,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
    txt(s,x+1.25,y+0.3,3.8,0.8,[(t,32,True,NAVY,HEAD)],anchor=MSO_ANCHOR.MIDDLE)
    txt(s,x+0.3,y+1.3,4.65,1.8,[(b,23,False,MUTE,BODY)])

# ---------- 3 tech stack ----------
s=S[2]
set_title(s,'Tech Stack & Architecture',80)
layers=[("CAPTURE","Flutter app (Android) · LiveKit WebRTC calling · Python agent hears caller and victim legs separately",TEAL),
        ("SPEECH","Sarvam STT & TTS: Hindi, Kannada, English and code-mixed speech, in and out",C(0x1F,0x6F,0xB5)),
        ("BRAIN","Keyword fast path + Jev typed judgments: calibrated, milliseconds per chunk · Gemini fallback",C(0x6A,0x4C,0xB5)),
        ("AUTOMATION","n8n (self-hosted): webhook trigger · AI Agent node with tools · Switch/IF branching · retries · human-approval step",RED),
        ("DATA & NOTIFY","Supabase: evidence log + shared scam-number list · Telegram alerts to trusted contacts",GREEN)]
for i,(lab,desc,col) in enumerate(layers):
    y=3.15+i*1.42
    box(s,1.125,y,17.75,1.25,fill=WHITE)
    box(s,1.125,y,3.6,1.25,fill=col,line=None)
    txt(s,1.125,y,3.6,1.25,[(lab,26,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
    txt(s,5.0,y,13.7,1.25,[(desc,22,False,NAVY,BODY)],anchor=MSO_ANCHOR.MIDDLE)
txt(s,1.125,10.3,17.75,0.6,[("Open, cheap and auditable: the classifier answers in milliseconds, and self-hosted n8n keeps call data under our control.",21,False,MUTE,HEAD)])

# ---------- 4 beyond automation ----------
s=S[3]
t=copy.deepcopy(S[2].shapes[[sh.name for sh in S[2].shapes].index('Google Shape;134;p3')]._element) if False else None
src=[sh for sh in S[1].shapes if sh.name=='Google Shape;127;p2'][0]
el=copy.deepcopy(src._element); s.shapes._spTree.append(el)
set_title(s,"Beyond Task Automation",80)
rows=[("Pressure ≥ 2 and asks for OTP / PIN","Whisper “Do not share that code”, ring the alarm, and ping trusted contacts now",RED),
      ("Poses as police / CBI and demands secrecy","Highest severity: send family the call summary and the next steps to follow",RED),
      ("Number already in the shared scam list","Raise confidence, block, and skip the guessing",AMBER),
      ("Call ends, victim confirms what happened","Agent drafts the 1930 / cybercrime.gov.in complaint in their language with transcript timestamps",TEAL),
      ("Money already moved","Switch to the time-critical checklist: bank freeze steps and helpline, in order",TEAL)]
txt(s,1.125,3.0,8,0.5,[("WHEN THE AGENT SEES…",20,True,TEAL,BODY)]); txt(s,10.2,3.0,8,0.5,[("…IT DECIDES AND ACTS",20,True,TEAL,BODY)])
for i,(a,b,col) in enumerate(rows):
    y=3.55+i*1.3
    box(s,1.125,y,8.6,1.15,fill=WHITE); box(s,1.125,y+0.1,0.12,0.95,fill=col,line=None,shape=MSO_SHAPE.RECTANGLE)
    txt(s,1.5,y,8.1,1.15,[(a,22,True,NAVY,BODY)],anchor=MSO_ANCHOR.MIDDLE)
    arrow(s,9.8,y+0.575,10.15,y+0.575,col)
    box(s,10.25,y,8.625,1.15,fill=WHITE)
    txt(s,10.5,y,8.2,1.15,[(b,20,False,MUTE,BODY)],anchor=MSO_ANCHOR.MIDDLE)
box(s,1.125,10.1,17.75,0.8,fill=NAVY,line=None)
txt(s,1.4,10.1,17.2,0.8,[("Multi-step reasoning, autonomous execution. A human confirms only the one irreversible step: submitting the complaint.",21,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE)

# ---------- 5 data flow ----------
s=S[4]
def node(x,y,w,h,title,sub,col=TEAL,fill=WHITE):
    box(s,x,y,w,h,fill=fill,line=col,lw=2.5)
    txt(s,x+0.2,y+0.1,w-0.4,h-0.2,[(title,25,True,col if col!=WHITE else NAVY,HEAD,5),(sub,17,False,MUTE,BODY)],anchor=MSO_ANCHOR.MIDDLE)
xs=[1.15,5.8,10.45,15.1]; W=3.7
node(xs[0],3.2,W,1.9,"Call audio","LiveKit WebRTC, caller & victim legs captured separately")
node(xs[1],3.2,W,1.9,"Sarvam STT","Indic speech to text, streaming chunks")
node(xs[2],3.2,W,1.9,"Jev classifier","Typed verdict: scam type, pressure 0–3, OTP ask, secrecy",C(0x6A,0x4C,0xB5))
node(xs[3],3.2,3.75,1.9,"Risk gate","Threshold crossed: both paths fire",RED)
for i in range(3): arrow(s,xs[i]+W+0.05,4.15,xs[i+1]-0.05,4.15)
arrow(s,16.97,5.15,16.97,6.0,RED)
txt(s,11.0,5.28,5.8,0.5,[("fires in parallel",17,False,RED,BODY)],align=PP_ALIGN.RIGHT)
arrow(s,16.97,5.62,8.5,5.62,RED,w=3,head=False); arrow(s,8.5,5.62,8.5,6.0,RED)
node(1.15,6.0,14.0,1.5,"n8n AI Agent (self-hosted)","Webhook receives the verdict + transcript · agent plans the response, calls tools, branches on severity",RED,fill=C(0xFD,0xF1,0xEF))
node(15.45,6.0,3.45,1.5,"Live path","Victim only",TEAL)
acts=[("Family alert","Summary + next steps to trusted contacts"),("1930 complaint","Draft in victim's language, confirm, file"),("Evidence + shared list","Supabase log, scam number added")]
for i,(a,b) in enumerate(acts):
    x=1.15+i*4.65; node(x,8.15,4.3,1.6,a,b,GREEN); arrow(s,x+2.15,7.55,x+2.25,8.1,GREEN)
node(14.95,8.15,3.95,1.6,"Alarm + TTS whisper","In their own language",TEAL); arrow(s,17.17,7.55,17.17,8.1,TEAL)
box(s,1.15,10.0,17.75,0.85,fill=NAVY,line=None)
txt(s,1.4,10.0,17.3,0.85,[("Feedback loop: every caught scam number joins the shared list and sharpens the next call's verdict.",21,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE)

# ---------- 6 screenshots (concept mockups) ----------
s=S[5]
set_title(s,"Screenshot of Your Project")
txt(s,1.125,2.75,17.7,0.5,[("Concept screens: the live build is demoed at the event.",20,False,MUTE,HEAD)])
def phone(x,y,w,h):
    box(s,x,y,w,h,fill=C(0x14,0x1B,0x22),line=None,r=0.07); return x+0.2,y+0.25,w-0.4,h-0.5
sx,sy,sw,sh_=phone(1.4,3.4,5.0,7.3)
box(s,sx,sy,sw,sh_,fill=C(0x3B,0x0D,0x0D),line=None,r=0.05)
txt(s,sx+0.3,sy+0.4,sw-0.6,6,[("⚠ LIVE CALL",18,True,C(0xFF,0xB4,0xAB),BODY,10),
 ("Possible digital-arrest scam",32,True,WHITE,HEAD,16),
 ("Caller claims to be CBI and asks you to keep this secret.",19,False,C(0xF3,0xD0,0xCB),BODY,16),
 ("PRESSURE  ■ ■ ■",20,True,C(0xFF,0x8A,0x7A),BODY,16),
 ("Do NOT share any OTP or PIN.",24,True,WHITE,BODY,6)])
b=box(s,sx+0.3,sy+sh_-1.2,sw-0.6,0.8,fill=C(0xE5,0x39,0x35),line=None,r=0.3)
txt(s,sx+0.3,sy+sh_-1.2,sw-0.6,0.8,[("END CALL",20,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
sx,sy,sw,sh_=phone(7.2,3.4,5.0,7.3)
box(s,sx,sy,sw,sh_,fill=WHITE,line=None,r=0.05)
txt(s,sx+0.3,sy+0.35,sw-0.6,6.5,[("Call summary",30,True,NAVY,HEAD,14)])
items=[("✔ Family alerted","2 trusted contacts",GREEN),("✔ 1930 complaint drafted","Review and confirm",GREEN),("✔ Evidence saved","Transcript + timestamps",GREEN),("✔ Number added","Shared scam list",GREEN)]
for i,(a,b_,col) in enumerate(items):
    yy=sy+1.3+i*1.3; box(s,sx+0.3,yy,sw-0.6,1.1,fill=C(0xEE,0xF8,0xF2),line=None,r=0.12)
    txt(s,sx+0.5,yy,sw-1.0,1.1,[(a,19,True,col,BODY,2),(b_,16,False,MUTE,BODY)],anchor=MSO_ANCHOR.MIDDLE)
# n8n canvas
cx,cy,cw,ch=13.0,3.4,5.9,7.3
box(s,cx,cy,cw,ch,fill=C(0x1E,0x20,0x2B),line=None,r=0.03)
txt(s,cx+0.3,cy+0.2,cw-0.6,0.5,[("n8n · Chaukanna response workflow",17,True,C(0xFF,0x6D,0x5A),BODY)])
def n8(x,y,label,col=C(0x3A,0x3E,0x52)):
    box(s,cx+x,cy+y,1.9,0.8,fill=col,line=C(0x8A,0x90,0xB0),lw=1.2,r=0.2)
    txt(s,cx+x,cy+y,1.9,0.8,[(label,14,True,WHITE,BODY)],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
n8(0.35,1.0,"Webhook"); n8(3.6,1.0,"AI Agent",C(0x7A,0x3E,0xD1)); n8(3.6,2.4,"Switch: severity")
n8(0.35,3.9,"Family alert",C(0x1E,0x7A,0x4E)); n8(3.6,3.9,"1930 draft",C(0x1E,0x7A,0x4E)); n8(0.35,5.4,"Wait: confirm",C(0xB7,0x7A,0x10)); n8(3.6,5.4,"Supabase log",C(0x1E,0x7A,0x4E))
for a in [(2.25,1.4,3.6,1.4),(4.55,1.8,4.55,2.4),(4.55,3.2,4.55,3.9),(3.6,3.2,1.3,3.9),(4.55,4.7,2.0,5.4),(4.55,4.7,4.55,5.4)]:
    arrow(s,cx+a[0],cy+a[1],cx+a[2],cy+a[3],C(0x8A,0x90,0xB0),w=2)
txt(s,1.4,10.8,5,0.4,[("Alarm screen",17,True,NAVY,BODY)],align=PP_ALIGN.CENTER); txt(s,7.2,10.8,5,0.4,[("After-call report",17,True,NAVY,BODY)],align=PP_ALIGN.CENTER); txt(s,13.0,10.8,5.9,0.4,[("n8n workflow",17,True,NAVY,BODY)],align=PP_ALIGN.CENTER)

# ---------- 7 others ----------
s=S[6]
set_title(s,"Others")
card(s,1.125,3.2,5.75,7.3,"Impact",["Protects the people scammers target most: elders and first-time digital users, in their own language.","Turns minutes of panic into steps that run on their own.","Every report grows a shared scam-number list that protects the next person."],accent=GREEN,tsz=38,bsz=27)
card(s,7.125,3.2,5.75,7.3,"Why n8n",["Visual, auditable response flow that judges and families can follow.","AI Agent node with tools gives multi-step reasoning, not just triggers.","Retries, branching and a human-approval step built in.","Self-hosted, so call data stays ours."],accent=RED,tsz=38,bsz=27)
card(s,13.125,3.2,5.75,7.3,"Roadmap",["Hackathon: app-to-app guardian calls with the full n8n response.","Next: real phone numbers via SIP (Exotel / Plivo / Twilio into LiveKit).","Then: more Indic languages, and sharing scam signals with banks and telecoms."],accent=TEAL,tsz=38,bsz=27)
p.save('Chaukanna_HackSprint_PS22_v2.pptx')
